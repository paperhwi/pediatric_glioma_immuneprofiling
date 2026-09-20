# -*- coding: utf-8 -*-
"""Build the two Word deliverables for the Cancers submission.

  Figure_and_Table_Legends_Cancers.docx
  Supplementary_Material_Cancers.docx   (figures + legends + table index)
"""
import os
import sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import pymupdf
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT

import legends as L
from build_tables import TABLES

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = f"{ROOT}/docs"
EMB = f"{ROOT}/panels/_embed"
os.makedirs(OUT, exist_ok=True)
os.makedirs(EMB, exist_ok=True)

BODY_PT = 10.5
MAX_W, MAX_H = 6.4, 7.6          # inches available on a Letter page


def style(doc):
    n = doc.styles["Normal"]
    n.font.name = "Calibri"
    n.font.size = Pt(BODY_PT)
    n.paragraph_format.space_after = Pt(8)
    n.paragraph_format.line_spacing = 1.15


def head(doc, text, size=13, space_before=14):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(size)
    return p


def legend_para(doc, label, title, body):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(11)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(f"{label}. ")
    r.bold = True
    r = p.add_run(title + " ")
    r.bold = True
    p.add_run(body)
    return p


def note(doc, text, italic=True, size=9.5):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    r.italic = italic
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor(0x40, 0x40, 0x40)
    return p


def embed_png(pdf_path, name, dpi_target=200):
    """Render a figure at the size it will occupy in the document."""
    d = pymupdf.open(pdf_path)
    page = d[0]
    w_in, h_in = page.rect.width / 72, page.rect.height / 72
    scale = min(MAX_W / w_in, MAX_H / h_in, 1.0)
    disp_w, disp_h = w_in * scale, h_in * scale
    dpi = min(dpi_target, 14000 / max(disp_w, disp_h))
    zoom = dpi * scale / 72
    pm = page.get_pixmap(matrix=pymupdf.Matrix(zoom, zoom), colorspace=pymupdf.csRGB)
    out = f"{EMB}/{name}.png"
    pm.save(out)
    d.close()
    return out, disp_w, disp_h


# ------------------------------------------------------------------ doc 1
def build_legends():
    doc = Document()
    style(doc)
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Inches(0.9)
        s.left_margin = s.right_margin = Inches(0.9)

    p = doc.add_paragraph()
    r = p.add_run("Figure and table legends")
    r.bold = True
    r.font.size = Pt(15)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    note(doc, L.TITLE, italic=True, size=11)
    note(doc, L.PREAMBLE)

    head(doc, "Main figures")
    for lab, title, body in L.MAIN:
        legend_para(doc, lab, title, body)

    head(doc, "Main tables")
    for lab, title, body in L.MAIN_TABLES:
        legend_para(doc, lab, title, body)

    head(doc, "Supplementary figures")
    for lab, title, body in L.SUPPL:
        legend_para(doc, lab, title, body)

    head(doc, "Supplementary tables")
    for num, title, desc, supports, _ in TABLES:
        legend_para(doc, f"Table {num}", title + ".", f"{desc} Supports: {supports}.")

    head(doc, "Abbreviations")
    doc.add_paragraph(L.ABBREV)

    head(doc, "Items deposited in the study repository")
    note(doc, "The following analyses were part of the earlier revision package and are not "
              "included in this submission. They are deposited, with the code and the executed "
              f"notebooks that produce them, at {L.GITHUB}.", italic=False, size=10)
    for lab, what in L.ARCHIVED:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(f"{lab} — ")
        r.bold = True
        p.add_run(what)

    path = f"{OUT}/Figure_and_Table_Legends_Cancers.docx"
    doc.save(path)
    return path


# ------------------------------------------------------------------ doc 2
def build_supplementary():
    doc = Document()
    style(doc)
    s = doc.sections[0]
    s.orientation = WD_ORIENT.PORTRAIT
    s.top_margin = s.bottom_margin = Inches(0.8)
    s.left_margin = s.right_margin = Inches(0.9)

    p = doc.add_paragraph()
    r = p.add_run("Supplementary Material")
    r.bold = True
    r.font.size = Pt(16)
    note(doc, L.TITLE, italic=True, size=11)
    note(doc, "Supplementary Figures S1–S16 and Supplementary Tables S1–S25. The tables are "
              "supplied as a separate workbook, Supplementary_Tables_Cancers.xlsx, whose README "
              "sheet repeats the index below and whose Crosswalk sheet maps every item of the "
              "previous version onto this one.")

    head(doc, "Contents", size=12)
    for lab, title, _ in L.SUPPL:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run(f"{lab}. ")
        r.bold = True
        r.font.size = Pt(9.5)
        r = p.add_run(title)
        r.font.size = Pt(9.5)
    for num, title, *_ in TABLES:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run(f"Table {num}. ")
        r.bold = True
        r.font.size = Pt(9.5)
        r = p.add_run(title)
        r.font.size = Pt(9.5)

    doc.add_page_break()
    head(doc, "Supplementary figures", size=13, space_before=0)
    for i, (lab, title, body) in enumerate(L.SUPPL, start=1):
        pdf = f"{ROOT}/suppl/FigureS{i}.pdf"
        png, w, h = embed_png(pdf, f"S{i}")
        pic = doc.add_paragraph()
        pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pic.paragraph_format.space_after = Pt(6)
        pic.add_run().add_picture(png, width=Inches(w))
        legend_para(doc, lab, title, body)
        if i < len(L.SUPPL):
            doc.add_page_break()

    doc.add_page_break()
    head(doc, "Supplementary tables", size=13, space_before=0)
    note(doc, "Supplied as Supplementary_Tables_Cancers.xlsx; one worksheet per table, named for "
              "the table number.")
    for num, title, desc, supports, _ in TABLES:
        legend_para(doc, f"Table {num}", title + ".", f"{desc} Supports: {supports}.")

    head(doc, "Items deposited in the study repository")
    note(doc, "The following analyses were part of the earlier revision package and are not "
              "included in this submission. They are deposited, with the code and the executed "
              f"notebooks that produce them, at {L.GITHUB}.", italic=False, size=10)
    for lab, what in L.ARCHIVED:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(f"{lab} — ")
        r.bold = True
        p.add_run(what)

    head(doc, "Abbreviations")
    doc.add_paragraph(L.ABBREV)

    path = f"{OUT}/Supplementary_Material_Cancers.docx"
    doc.save(path)
    return path


if __name__ == "__main__":
    for f in (build_legends(), build_supplementary()):
        print(f"{os.path.basename(f):<44} {os.path.getsize(f)/1e6:.2f} MB")
