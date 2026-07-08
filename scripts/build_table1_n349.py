"""
Table 1 — Clinical and molecular characteristics of the Main cohort (n=349).

Columns: DMG_K27 (n=175) · DHG_G34 (n=31) · pHGG_WT (n=125) · IHG (n=18) · Total (n=349)
Rows  : Age, Sex, Location, Molecular subtype, BRAF/RTK fusion, OS availability, Ecotype.

Outputs:
  - output/Table1_n349.tsv  (machine-readable)
  - output/Table1_n349.docx (Word document with proper formatting)
"""
from __future__ import annotations
from pathlib import Path
import numpy as np, pandas as pd
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = Path("/sessions/zealous-vibrant-rubin/mnt/Open PBTA")
OUT  = BASE / "output"

main = pd.read_csv(OUT / "cohort_main_final.tsv", sep="\t").set_index("Kids_First_Biospecimen_ID")
ec3  = pd.read_csv(OUT / "ecotype_LM22_main_k3_annotated.tsv", sep="\t").set_index("Kids_First_Biospecimen_ID")
df = main.join(ec3[["ecotype"]], how="left")

cohort_order = ["DMG_K27", "DHG_G34", "pHGG_WT", "IHG"]
cohort_labels = {"DMG_K27":"DMG H3 K27-altered", "DHG_G34":"DHG H3 G34-mutant",
                  "pHGG_WT":"pHGG H3/IDH-wildtype", "IHG":"Infant-type hemispheric glioma"}
groups = {c: df[df["cohort_group"] == c] for c in cohort_order}
groups["Total"] = df
labels_in_order = cohort_order + ["Total"]

def fmt_pct(n, denom):
    return f"{int(n)} ({100*n/denom:.0f}%)" if denom else "0 (0%)"

def median_iqr(s):
    s = pd.to_numeric(s, errors="coerce").dropna()
    if len(s) == 0: return "—"
    return f"{s.median():.1f} [{s.quantile(0.25):.1f}–{s.quantile(0.75):.1f}]"

def count_in(s, value):
    if pd.isna(value):
        return int(s.isna().sum())
    return int((s == value).sum())

def collect_levels(col, drop_na=True, top_k=None):
    vc = df[col].value_counts(dropna=drop_na)
    if top_k is not None:
        return vc.head(top_k).index.tolist()
    return vc.index.tolist()

# ---------------------------------------------------------------------------
# Build rows: each row is a (label, fn(group_df) -> str)
# ---------------------------------------------------------------------------
rows = []

# Section header markers will be (label, None)
def header(s): rows.append((s, None, True))
def value(label, fn): rows.append((label, fn, False))

# Sample count
header("Samples, n")
value("Independent-primary biospecimens", lambda g: f"{len(g)}")
value("Unique patients", lambda g: f"{g['Kids_First_Participant_ID'].nunique()}")

# Age
header("Age at diagnosis, years")
value("Median [IQR]", lambda g: median_iqr(g["age_years"]))
value("Range", lambda g: (
    f"{pd.to_numeric(g['age_years'], errors='coerce').min():.1f}–"
    f"{pd.to_numeric(g['age_years'], errors='coerce').max():.1f}"
    if pd.to_numeric(g['age_years'], errors='coerce').notna().any() else "—"))
value("Age (n missing)", lambda g: f"{int(pd.to_numeric(g['age_years'], errors='coerce').isna().sum())}")
header("Developmental age group, n (%)")
for grp_label, grp_value in [("Infant (<1 yr)", "infant"),
                              ("Young child (1–4 yr)", "young_child"),
                              ("Child (5–11 yr)", "child"),
                              ("Adolescent (≥12 yr)", "adolescent")]:
    value("  "+grp_label, lambda g, v=grp_value: fmt_pct((g["age_dev_group"]==v).sum(), len(g)))

# Sex
header("Sex, n (%)")
value("  Male",   lambda g: fmt_pct((g["reported_gender"]=="Male").sum(), len(g)))
value("  Female", lambda g: fmt_pct((g["reported_gender"]=="Female").sum(), len(g)))
value("  Unknown/Not reported",
      lambda g: fmt_pct(((g["reported_gender"]=="Unknown")|
                          (g["reported_gender"]=="Not Reported")|
                          g["reported_gender"].isna()).sum(), len(g)))

# Anatomical compartment
header("Tumor location, n (%)")
value("  Midline",        lambda g: fmt_pct((g["location_class"]=="Midline").sum(),         len(g)))
value("  Hemispheric",    lambda g: fmt_pct((g["location_class"]=="Hemispheric").sum(),     len(g)))
value("  Posterior fossa",lambda g: fmt_pct((g["location_class"]=="Posterior_fossa").sum(), len(g)))
value("  Ambiguous (pineal/suprasellar/etc.)",
                          lambda g: fmt_pct((g["location_class"]=="Ambiguous").sum(),       len(g)))
value("  Other / not annotated",
                          lambda g: fmt_pct((g["location_class"]=="Other/NA").sum(),        len(g)))

# Molecular subtype
header("Molecular subtype, n (%)")
def molsub_count(g, prefix, with_tp53=False):
    vals = g["molecular_subtype"].dropna().astype(str)
    if with_tp53:
        n = vals.str.contains(prefix).sum() & vals.str.contains("TP53").sum()
        n = int((vals.str.contains(prefix) & vals.str.contains("TP53")).sum())
    else:
        n = int(vals.str.contains(prefix).sum())
    return fmt_pct(n, len(g))

value("  H3 K28-altered (formerly K27-altered)",
      lambda g: fmt_pct(int(g["molecular_subtype"].astype(str).str.contains("H3 K28").sum()), len(g)))
value("    with TP53 alteration",
      lambda g: fmt_pct(int(((g["molecular_subtype"].astype(str).str.contains("H3 K28")) &
                              (g["molecular_subtype"].astype(str).str.contains("TP53"))).sum()), len(g)))
value("  H3 G35-mutant (formerly G34-mutant)",
      lambda g: fmt_pct(int(g["molecular_subtype"].astype(str).str.contains("H3 G35").sum()), len(g)))
value("    with TP53 alteration",
      lambda g: fmt_pct(int(((g["molecular_subtype"].astype(str).str.contains("H3 G35")) &
                              (g["molecular_subtype"].astype(str).str.contains("TP53"))).sum()), len(g)))
value("  H3/IDH wildtype HGG",
      lambda g: fmt_pct(int(g["molecular_subtype"].astype(str).str.contains("H3 wildtype").sum()), len(g)))
value("    with TP53 alteration",
      lambda g: fmt_pct(int(((g["molecular_subtype"].astype(str).str.contains("H3 wildtype")) &
                              (g["molecular_subtype"].astype(str).str.contains("TP53"))).sum()), len(g)))
value("  Infant-type, NTRK/ROS1/MET/ALK-altered",
      lambda g: fmt_pct(int(g["molecular_subtype"].astype(str).str.contains(
          "IHG|NTRK|ROS1|MET-|ALK-", regex=True).sum()), len(g)))
value("  Other / to be classified",
      lambda g: fmt_pct(int(g["molecular_subtype"].astype(str).str.contains(
          "To be classified|GNG|GNT", regex=True).sum()), len(g)))

# BRAF / IHG fusion flags
header("Driver alteration flags, n")
value("  has_braf_fusion (any)", lambda g: f"{int(g['has_braf_fusion'].sum())}")
value("  has_ihg_fusion (NTRK/ROS1/MET/ALK)",  lambda g: f"{int(g['has_ihg_fusion'].sum())}")

# Survival availability
header("Survival data availability, n (%)")
value("  OS days available",
      lambda g: fmt_pct(int(pd.to_numeric(g["OS_days"], errors="coerce").notna().sum()), len(g)))
value("  Events (deceased)",
      lambda g: fmt_pct(int((g["OS_status"].astype(str).str.upper()=="DECEASED").sum()), len(g)))

# Ecotype distribution
header("Immune ecotype (this study, k=3), n (%)")
for ec in ["Lymphocyte-inflamed", "Myeloid-dominant", "Immune-desert"]:
    value(f"  {ec}",
          lambda g, e=ec: fmt_pct(int((g["ecotype"]==e).sum()), len(g)))
value("  Not assigned / missing",
      lambda g: fmt_pct(int(g["ecotype"].isna().sum()), len(g)))

# ---------------------------------------------------------------------------
# Build dataframe
# ---------------------------------------------------------------------------
out_rows = []
for entry in rows:
    label = entry[0]; fn = entry[1]; is_header = entry[2]
    if is_header:
        out_rows.append({"Variable": label, **{c:"" for c in labels_in_order}})
    else:
        row = {"Variable": label}
        for c in labels_in_order:
            row[c] = fn(groups[c])
        out_rows.append(row)
table = pd.DataFrame(out_rows)

# Column headers with counts
col_label = {c: f"{cohort_labels.get(c, c)} (n={len(groups[c])})" if c != "Total" else f"Total (n={len(df)})"
             for c in labels_in_order}
table = table.rename(columns=col_label)
table.to_csv(OUT / "Table1_n349.tsv", sep="\t", index=False)
print(f"Wrote: Table1_n349.tsv  ({len(table)} rows)")

# ---------------------------------------------------------------------------
# Build Word docx
# ---------------------------------------------------------------------------
doc = Document()
section = doc.sections[0]
section.top_margin = Cm(2.0); section.bottom_margin = Cm(2.0)
section.left_margin = Cm(1.8); section.right_margin = Cm(1.8)

# Title
title = doc.add_paragraph()
run = title.add_run("Table 1.  Clinical and molecular characteristics of the pediatric glioma Main cohort (n = 349).")
run.font.size = Pt(11); run.bold = True

# Table
ncols = len(table.columns)
tbl = doc.add_table(rows=len(table)+1, cols=ncols)
tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
tbl.autofit = True

# Header row
hdr = tbl.rows[0]
for i, c in enumerate(table.columns):
    cell = hdr.cells[i]
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(c)
    r.bold = True; r.font.size = Pt(9)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i > 0 else WD_ALIGN_PARAGRAPH.LEFT
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

# Data rows
for ri, (_, row) in enumerate(table.iterrows(), start=1):
    is_header_row = all(v == "" for v in row.values[1:])
    for ci, col in enumerate(table.columns):
        cell = tbl.rows[ri].cells[ci]
        cell.text = ""
        p = cell.paragraphs[0]
        r = p.add_run(str(row[col]))
        r.font.size = Pt(9)
        if is_header_row and ci == 0:
            r.bold = True; r.italic = True
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if ci > 0 else WD_ALIGN_PARAGRAPH.LEFT
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

# Borders
def set_cell_border(cell, **kwargs):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for side in ("top","left","bottom","right"):
        bd = OxmlElement(f"w:{side}")
        bd.set(qn("w:val"), "single"); bd.set(qn("w:sz"), "4"); bd.set(qn("w:color"), "000000")
        tcBorders.append(bd)
    tcPr.append(tcBorders)
for row in tbl.rows:
    for cell in row.cells:
        set_cell_border(cell)

# Footnote
foot = doc.add_paragraph()
fr = foot.add_run(
    "DMG, diffuse midline glioma; DHG, diffuse hemispheric glioma; pHGG, pediatric-type high-grade glioma; "
    "IHG, infant-type hemispheric glioma. Median age reported as median [IQR]. "
    "Ecotype assignment from consensus clustering on LM22 + brain-tuned ssGSEA features (k=3; "
    "Methods §2). 'Not assigned / missing' indicates samples without ecotype call (feature NA after z-scoring)."
)
fr.font.size = Pt(8); fr.italic = True

out_docx = OUT / "Table1_n349.docx"
doc.save(out_docx)
print(f"Wrote: {out_docx}")
