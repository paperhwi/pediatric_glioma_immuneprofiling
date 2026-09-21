import pandas as pd, numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

ARIAL=lambda **k: Font(name="Arial", **k)
HDR_FILL=PatternFill("solid", fgColor="E8EDF5")
TITLE=ARIAL(size=12,bold=True,color="1E3A5F")
THIN=Side(style="thin",color="BFBFBF")

# module gene lists
WANT={"Bone Marrow-L2-Hematopoeitic Stem Cell":"HSC",
      "Bone Marrow-L2-Lymphoid Primed Multipotent Progenitor":"LMPP",
      "Bone Marrow-L2-Granulocyte Monocyte Progenitor":"GMP",
      "Bone Marrow-L2-Common Lymphoid Progenitor":"CLP",
      "Bone Marrow-L2-Erythroid Megakaryocyte Progenitor":"EMP (lineage control)",
      "Bone Marrow-L2-CD14 Monocyte":"CD14 monocyte (positive control)",
      "Bone Marrow-L2-Macrophage":"Macrophage (positive control)"}
sets={}
for line in open("Azimuth_2023.gmt"):
    f=line.rstrip("\n").split("\t")
    if f[0] in WANT: sets[WANT[f[0]]]=[g.split(",")[0] for g in f[2:] if g.strip()]
gl=pd.DataFrame([dict(module=k,n_genes=len(v),genes=", ".join(v),
                      source="Azimuth 2023 human bone-marrow reference (Hao et al. 2021)") for k,v in sets.items()])

hspc_kw=pd.read_csv("T4_KW_by_ecotype.tsv",sep="\t")
hspc_dunn=pd.read_csv("T4_Dunn_posthoc.tsv",sep="\t")
hspc_res=pd.read_csv("T4_residualised_specificity.tsv",sep="\t")
hspc_stats=pd.concat([
    hspc_kw.assign(analysis="Kruskal-Wallis across ecotypes"),
    hspc_res.assign(analysis="Residualised on mature monocyte + macrophage content"),
],ignore_index=True)

TABLES=[
 ("S12","Cluster-number sweep, k = 2 to 10",
  "Consensus clustering of the primary 46-feature matrix (n = 349). B = 1,000 bootstrap replicates, 80% subsampling, "
  "KMeans n_init = 1, seed 42; final labels by average linkage on the 1 - consensus distance. PAC is the fraction of "
  "off-diagonal consensus values in (0.1, 0.9). delta_AUC is the relative increase in the area under the consensus "
  "cumulative-distribution curve over the preceding k and is undefined at k = 2. Values for k = 2 to 6 reproduce the "
  "locked pipeline output to < 5e-5.",
  pd.read_csv("T1_ksweep_k2_k10.tsv",sep="\t")),
 ("S13","Gap statistic, k = 1 to 10",
  "Tibshirani (2001) gap statistic against 50 uniform reference datasets generated over the bounding box of the "
  "principal-component rotation of the observed data, seed 42. meets_1SE_rule is TRUE where gap(k) >= gap(k+1) - s(k+1); "
  "the smallest such k is the selected number of clusters. This is the only internal index defined at k = 1.",
  pd.read_csv("T1_gap_statistic.tsv",sep="\t")),
 ("S14","Cluster-number criteria and the k each selects",
  "Criteria are grouped by what they measure. Separation criteria assess how well separated a partition is at a given k; "
  "number-of-clusters criteria compare the observed structure against a null reference. The two families disagree, which "
  "is reported in the manuscript rather than resolved by preferring one.",
  pd.read_csv("T1_criterion_summary.tsv",sep="\t")),
 ("S15","Overlap of the three adjusted differential-expression contrasts",
  "Genes at BH q < 0.05 and molecular-group-adjusted log2 fold change > 1. Both Immune-desert-referenced contrasts share "
  "the immune-presence axis; the Lymphocyte-inflamed versus Myeloid-dominant contrast is the one that distinguishes the "
  "two immune-present ecotypes.",
  pd.read_csv("T2_contrast_overlap_summary.tsv",sep="\t")),
 ("S16","Genes shared by both Immune-desert-referenced contrasts",
  "The 426 genes up-regulated at BH q < 0.05 and adjusted log2 fold change > 1 in both Lymphocyte-inflamed versus "
  "Immune-desert and Myeloid-dominant versus Immune-desert, with their adjusted log2 fold change in all three contrasts. "
  "The Lymphocyte-inflamed effect is a median 1.61-fold larger (paired Wilcoxon P = 7.6e-71).",
  pd.read_csv("T2_shared_immune_presence_genes.tsv",sep="\t")),
 ("S17","Analysis-set reconciliation",
  "The location analyses and the survival analyses draw different subsets of the same 349 tumours and are not nested. "
  "Of the 251 survival-evaluable tumours, 180 fall within the n = 258 sequential-PERMANOVA subset and 3 have no usable "
  "location annotation. Referenced by Figure 1B.",
  pd.read_csv("T5_analysis_set_reconciliation.tsv",sep="\t")),
 ("S18","Survival-evaluability by baseline characteristic",
  "Comparison of the 251 tumours with overall-survival annotation against the 98 without. Survival-evaluability is not "
  "associated with immune ecotype (P = 0.463) but is associated with integrated molecular group and anatomical location, "
  "both of which are covariates in the multivariable Cox model and in the inverse-probability-of-observation weighting model.",
  pd.read_csv("T3_SupplTable_OS_evaluability.tsv",sep="\t")),
 ("S19","Ecotype and survival-evaluability within each molecular group",
  "Stratified test removing confounding by molecular group. The association remains non-significant in every group except "
  "infant-type hemispheric glioma, where 18 tumours give cell counts too small to interpret.",
  pd.read_csv("T3_within_group_ecotype_x_OSavail.tsv",sep="\t")),
 ("S20","Haematopoietic stem and progenitor module gene lists",
  "Marker gene sets for the bone-marrow cell types scored in Supplementary Figure S19, retrieved from the Azimuth 2023 "
  "human bone-marrow reference via Enrichr on 16 September 2026. The progenitor modules share no genes with the mature "
  "monocyte or macrophage modules.",
  gl),
 ("S21","Haematopoietic stem and progenitor ssGSEA scores",
  "Per-sample single-sample gene set enrichment scores (gseapy, sample_norm_method = 'rank') on log2(TPM + 1), n = 349, "
  "with the locked ecotype assignment.",
  pd.read_csv("T4_HSPC_ssGSEA_scores.tsv",sep="\t")),
 ("S22","Haematopoietic stem and progenitor programs by ecotype",
  "Kruskal-Wallis tests across the three ecotypes, and the same tests after residualising each module on mature CD14 "
  "monocyte and macrophage content by ordinary least squares. No progenitor module remains associated with ecotype after "
  "adjustment (all BH q >= 0.18); the granulocyte-monocyte progenitor effect falls from epsilon-squared 0.205 to 0.007.",
  hspc_stats),
 ("S23","Dunn post-hoc comparisons for the progenitor modules",
  "Pairwise Dunn tests with Benjamini-Hochberg correction across all 21 comparisons. Haematopoietic stem-cell scores do "
  "not differ between Lymphocyte-inflamed and Myeloid-dominant (z = 0.29, P = 0.775).",
  hspc_dunn),
 ("S24","Software versions, parameters and random seeds",
  "Execution environment for the analyses added in this revision. The inherited pipeline environment is reported "
  "separately in the existing reproducibility record.",
  pd.read_csv("T5_software_environment.tsv",sep="\t")),
]

wb=Workbook(); ws=wb.active; ws.title="README"
ws["A1"]="Supplementary tables S12-S24 added in revision"; ws["A1"].font=ARIAL(size=14,bold=True,color="1E3A5F")
ws["A2"]="Cancer Informatics · CIX-26-0214 · 16 September 2026"; ws["A2"].font=ARIAL(size=10,color="595959")
ws["A3"]="Integrated Transcriptomic Immune Profiling Identifies Survival-Associated Immune Ecotypes in Pediatric Diffuse High-Grade Glioma"
ws["A3"].font=ARIAL(size=10,italic=True,color="595959")
r=5
for h,w in zip(["Table","Sheet","Title","Addresses","Rows"],[10,10,62,16,9]):
    c=ws.cell(row=r,column=["Table","Sheet","Title","Addresses","Rows"].index(h)+1,value=h)
    c.font=ARIAL(size=10,bold=True); c.fill=HDR_FILL
    ws.column_dimensions[get_column_letter(["Table","Sheet","Title","Addresses","Rows"].index(h)+1)].width=w
ADDR={"S12":"R1-2, R2-1","S13":"R1-2","S14":"R1-2","S15":"R1-1","S16":"R1-1","S17":"R2 minor",
      "S18":"R1-6","S19":"R1-6","S20":"R1-7","S21":"R1-7","S22":"R1-7","S23":"R1-7","S24":"R2 minor, Editor"}
r=6
for tid,title,desc,df in TABLES:
    for j,v in enumerate([f"Table {tid}",tid,title,ADDR[tid],len(df)],start=1):
        c=ws.cell(row=r,column=j,value=v); c.font=ARIAL(size=10)
        c.alignment=Alignment(vertical="top",wrap_text=(j==3))
    r+=1
ws.cell(row=r+1,column=1,value="Every table is regenerated by the notebooks in 04_notebooks/ of the deposited repository; each notebook reproduces the previously published values before extending them.").font=ARIAL(size=9,italic=True,color="595959")
ws.freeze_panes="A6"

for tid,title,desc,df in TABLES:
    s=wb.create_sheet(tid)
    s["A1"]=f"Supplementary Table {tid}. {title}"; s["A1"].font=TITLE
    s["A2"]=desc; s["A2"].font=ARIAL(size=9,color="404040"); s["A2"].alignment=Alignment(wrap_text=True,vertical="top")
    s.merge_cells(start_row=2,start_column=1,end_row=4,end_column=max(6,min(len(df.columns),12)))
    hr=6
    for j,col in enumerate(df.columns,start=1):
        c=s.cell(row=hr,column=j,value=str(col)); c.font=ARIAL(size=10,bold=True); c.fill=HDR_FILL
        c.border=Border(bottom=THIN); c.alignment=Alignment(vertical="bottom",wrap_text=True)
    for i,(_,row) in enumerate(df.iterrows(),start=hr+1):
        for j,col in enumerate(df.columns,start=1):
            v=row[col]
            if isinstance(v,(np.integer,)): v=int(v)
            elif isinstance(v,(np.floating,)): v=None if pd.isna(v) else float(v)
            elif isinstance(v,(np.bool_,)): v=bool(v)
            elif pd.isna(v): v=None
            else: v=v if isinstance(v,(int,float,bool)) else str(v)
            c=s.cell(row=i,column=j,value=v); c.font=ARIAL(size=10)
            if isinstance(v,float) and abs(v)<1e-3 and v!=0: c.number_format="0.00E+00"
    for j,col in enumerate(df.columns,start=1):
        L=max(len(str(col)),*(len(str(x)) for x in df[col].head(200)))
        s.column_dimensions[get_column_letter(j)].width=min(max(L+2,10),58)
    s.freeze_panes=f"A{hr+1}"

wb.save("/mnt/user-data/outputs/w1/01_tables/Supplementary_tables_S12-S24_revision.xlsx")
print("saved; sheets:",len(wb.sheetnames),wb.sheetnames)
