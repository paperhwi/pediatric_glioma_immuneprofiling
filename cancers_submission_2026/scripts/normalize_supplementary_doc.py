"""Restore all supplementary labels and cross-references from their source lists."""
from pathlib import Path
import sys
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(Path(__file__).parent))
import legends as L
from build_tables import TABLES

PATH = ROOT / "4. Supplementary material" / "Supplementary_Material_Cancers.docx"
doc = Document(PATH)

# The initial contents list has one compact paragraph per supplementary figure.
# Reinsert S5 if a prior removal of the old S6 also removed its adjacent list item.
if not any(p.text.startswith("Figure S5.") for p in doc.paragraphs[:30]):
    for paragraph in doc.paragraphs[:30]:
        if paragraph.text.startswith("Figure S6."):
            entry = paragraph.insert_paragraph_before(style=paragraph.style)
            label, title, _ = L.SUPPL[4]
            entry.paragraph_format.space_after = paragraph.paragraph_format.space_after
            entry.add_run(f"{label}. ").bold = True
            entry.add_run(title)
            break

for label, title, body in L.SUPPL:
    for paragraph in doc.paragraphs:
        if title not in paragraph.text or not paragraph.text.startswith("Figure S"):
            continue
        runs = paragraph.runs
        if len(runs) < 2:
            continue
        runs[0].text = f"{label}. "
        runs[1].text = f"{title} " if len(runs) >= 3 else title
        if len(runs) >= 3:
            runs[2].text = body

for number, title, desc, supports, _ in TABLES:
    for paragraph in doc.paragraphs:
        if not paragraph.text.startswith(f"Table {number}.") or title not in paragraph.text:
            continue
        runs = paragraph.runs
        if len(runs) >= 3:
            runs[2].text = f"{desc} Supports: {supports}."

doc.save(PATH)
print(PATH)
