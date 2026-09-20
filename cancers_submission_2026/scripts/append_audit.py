# -*- coding: utf-8 -*-
"""Append the citation checklist to the supplementary table workbook."""
import os
import sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import openpyxl
from openpyxl.styles import Font
import verify

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
WB = f"{ROOT}/tables/Supplementary_Tables_Cancers.xlsx"

SECTION = {
    "Figure 1": "Results, cohort overview; Methods, cohort assembly",
    "Figure 2": "Results, deconvolution and cross-method benchmarking",
    "Figure 3": "Results, consensus clustering and cluster-number selection",
    "Figure 4": "Results, molecular group and anatomical location",
    "Figure 5": "Results, single-cell and spatial support",
    "Figure 6": "Results, differential expression",
    "Figure 7": "Results, survival",
    "Figure S1": "Results, deconvolution",
    "Figure S2": "Results, deconvolution",
    "Figure S3": "Methods, gene-expression processing",
    "Figure S4": "Results, consensus clustering",
    "Figure S5": "Results, sensitivity analyses",
    "Figure S6": "Results, molecular group and anatomical location",
    "Figure S7": "Results, immune programs across the ecotypes",
    "Figure S8": "Results, immune programs across the ecotypes",
    "Figure S9": "Results, differential expression",
    "Figure S10": "Results, differential expression",
    "Figure S11": "Results, subgroup survival",
    "Figure S12": "Results, subgroup survival",
    "Figure S13": "Results, survival (proportional-hazards assumption)",
    "Figure S14": "Results, selection bias and sensitivity analyses",
    "Figure S15": "Results, single-cell and spatial support",
    "Figure S16": "Results, single-cell and spatial support",
}

HEAD = Font(bold=True, size=10)


def main():
    rows, problems = verify.main()
    wb = openpyxl.load_workbook(WB)
    if "Citation_check" in wb.sheetnames:
        wb.remove(wb["Citation_check"])
    ws = wb.create_sheet("Citation_check")
    ws["A1"] = "Citation checklist for the submitted display items"
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = ("Every item below needs at least one citation in the running text. The fourth "
                "column records where it is already cross-referenced from another legend, which "
                "does not replace a main-text citation.")
    ws["A2"].font = Font(italic=True, size=10, color="404040")
    for j, h in enumerate(["Type", "Item", "Title", "Cross-referenced from",
                           "Suggested section for the main-text citation",
                           "Main-text citation"], start=1):
        ws.cell(4, j, h).font = HEAD
    for i, (typ, item, title, xref, need) in enumerate(rows, start=5):
        ws.cell(i, 1, typ)
        ws.cell(i, 2, item)
        ws.cell(i, 3, title)
        ws.cell(i, 4, xref)
        ws.cell(i, 5, SECTION.get(item, ""))
        ws.cell(i, 6, need)
    for col, w in zip("ABCDEF", [21, 13, 62, 40, 50, 20]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A5"
    # keep it next to the crosswalk
    order = ["README", "Crosswalk", "Citation_check"]
    wb._sheets = [wb[n] for n in order] + [s for s in wb._sheets if s.title not in order]
    wb.save(WB)
    print(f"\nappended Citation_check ({len(rows)} items) to {os.path.basename(WB)}")
    return problems


if __name__ == "__main__":
    main()
