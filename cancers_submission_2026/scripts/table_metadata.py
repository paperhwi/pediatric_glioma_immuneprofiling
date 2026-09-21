"""Read final supplementary-table metadata from the submitted workbook."""

from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
CANDIDATES = (
    ROOT / "supplementary_material" / "Supplementary_Tables_Cancers.xlsx",
    ROOT / "4. Supplementary material" / "Supplementary_Tables_Cancers.xlsx",
)


def load_table_metadata():
    workbook_path = next((path for path in CANDIDATES if path.exists()), None)
    if workbook_path is None:
        raise FileNotFoundError("Supplementary_Tables_Cancers.xlsx was not found")
    workbook = load_workbook(workbook_path, read_only=True, data_only=False)
    rows = []
    for row in workbook["README"].iter_rows(min_row=6, values_only=True):
        number, title, supports, _, description = row[:5]
        if isinstance(number, str) and number.startswith("S") and number[1:].isdigit():
            rows.append((number, title, description, supports, None))
    return rows


TABLES = load_table_metadata()
